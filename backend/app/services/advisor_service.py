from __future__ import annotations

import os
from copy import deepcopy
from pathlib import Path

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session

from app.models.advisor import AdvisorAnswer
from app.schemas.advisor import AdvisorRead


CX_WORKBOOK_FILENAME = 'CX Professional Services Tool Matrix_Jul 13, 2026.xlsx'
CX_WORKBOOK_ENV_VAR = 'CX_TOOL_MATRIX_WORKBOOK_PATH'
CX_WORKBOOK_REPO_PATH = Path(__file__).resolve().parents[1] / 'data' / CX_WORKBOOK_FILENAME


def _resolve_cx_workbook_path() -> Path | None:
    override = os.environ.get(CX_WORKBOOK_ENV_VAR)
    if override:
        path = Path(override)
        return path if path.exists() else None

    if CX_WORKBOOK_REPO_PATH.exists():
        return CX_WORKBOOK_REPO_PATH

    default_path = Path.home() / 'Downloads' / CX_WORKBOOK_FILENAME
    if default_path.exists():
        return default_path

    return None


def _clean_workbook_value(value: object) -> object:
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def _node_reference(target: object) -> str | None:
    target = _clean_workbook_value(target)
    if target is None:
        return None
    if isinstance(target, int):
        if 1 <= target <= 12:
            return f'cx-{target}'
        return f'cx-o{target}'
    if isinstance(target, str) and target.lower() == 'end':
        return 'cx-o20'
    return None


def _apply_cx_workbook_overrides(advisor: dict[str, object]) -> None:
    workbook_path = _resolve_cx_workbook_path()
    if workbook_path is None:
        return

    workbook = load_workbook(workbook_path, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    advisor['name'] = str(worksheet.cell(2, 1).value or advisor.get('name', 'CX Professional Services Tool Matrix'))
    advisor['description_html'] = worksheet.cell(2, 2).value or advisor.get('description_html')
    advisor['max_questions'] = 12

    question_lookup = {question['id']: question for question in advisor.get('questions', [])}
    outcome_lookup = {outcome['id']: outcome for outcome in advisor.get('outcomes', [])}

    for row_index in range(4, worksheet.max_row + 1):
        node = _clean_workbook_value(worksheet.cell(row_index, 1).value)
        if node is None or isinstance(node, str) and node == 'End':
            continue

        question_text = worksheet.cell(row_index, 2).value
        guidance_text = worksheet.cell(row_index, 3).value
        yes_target = _node_reference(worksheet.cell(row_index, 4).value)
        no_target = _node_reference(worksheet.cell(row_index, 5).value)
        details_html = worksheet.cell(row_index, 7).value

        if isinstance(node, int) and node <= 12:
            question = question_lookup.get(f'cx-{node}')
            if question is None:
                continue
            question['prompt'] = str(question_text or question.get('prompt', ''))
            question['guidance'] = str(guidance_text or question.get('guidance', ''))
            question['yes_next_question_id'] = yes_target if yes_target and yes_target.startswith('cx-') and not yes_target.startswith('cx-o') else None
            question['no_next_question_id'] = no_target if no_target and no_target.startswith('cx-') and not no_target.startswith('cx-o') else None
            question['yes_outcome_id'] = yes_target if yes_target and yes_target.startswith('cx-o') else None
            question['no_outcome_id'] = no_target if no_target and no_target.startswith('cx-o') else None
            continue

        if isinstance(node, int):
            outcome = outcome_lookup.get(f'cx-o{node}')
            if outcome is None:
                continue
            if details_html is not None:
                outcome['details_html'] = details_html



ADVISOR_CATALOG: list[dict[str, object]] = [
    {
        'id': 'cx-professional-services-tool-matrix',
        'name': 'CX Professional Services Tool Matrix',
        'description': 'Guidance on tools for storing and sharing CX Professional Services documents.',
        'description_html': (
            '<p><span>Welcome Professional Services Tools Matrix users! </span></p>'
            '<p>Below you will see a new <span>advisor that will</span> <span>provide the same '
            'guidance on what tools you should use to store and share CX Professional Services '
            'Documents. This advisor asks a few questions about the files that you need to store or '
            'share in order to point you towards the best tool to use. Please answer each question '
            'as best you can in order to be pointed towards the most appropriate tool. In Data '
            'Advisor, you may also explore features such as <strong><a href="https://dataadvisor.cloudapps.cisco.com/da/tools">Tool Advisor</a></strong>, '
            '<strong><a href="https://dataadvisor.cloudapps.cisco.com/da/cwizard">Data Classification Wizard</a></strong>, '
            '<strong><a href="https://dataadvisor.cloudapps.cisco.com/da/categories">Data Categories</a></strong> and '
            '<strong><a href="https://dataadvisor.cloudapps.cisco.com/da/advisors">More Advisors</a></strong>. '
            'We’d also ask that you visit the <strong><a href="https://dataadvisor.cloudapps.cisco.com/da/tools">Do Not Use Tools page</a></strong> '
            'in Data Advisor to make sure you stay in compliance.</span></p>'
            '<p><span style="color:#ff0000;"><strong>Note:</strong></span> If you have any questions, '
            'please contact the <strong><a href="https://pwc018.cloudapps.cisco.com/wsrp/pwc018/mobile/view/widget/294054/InTouch">Data Advisor Team</a></strong>.</p>'
        ),
        'max_questions': 12,
        'questions': [
            {
                'id': 'cx-1',
                'prompt': (
                    'Are you looking for information regarding Customer Experience (CX) Professional '
                    'Services Customer Delivery Projects ?'
                ),
                'guidance': (
                    'This advisor provides guidance on management of CX Professional Services '
                    'Customer Delivery Projects documentation'
                ),
                'answer': None,
                'yes_next_question_id': 'cx-2',
                'no_outcome_id': 'cx-o20',
            },
            {
                'id': 'cx-2',
                'prompt': 'Are you looking for information on storage and sharing of project documents ?',
                'guidance': 'Management of CX Professional Services Customer Delivery Projects documentation is subject to the Customer Experience Professional Services Document Management Policy',
                'answer': None,
                'yes_next_question_id': 'cx-3',
                'no_outcome_id': 'cx-o20',
            },
            {
                'id': 'cx-3',
                'prompt': 'Are you looking for information on storage and sharing of Deliverables ?',
                'guidance': 'Management of Deliverables must be via approved methods as described in this section.',
                'answer': None,
                'yes_next_question_id': 'cx-6',
                'no_next_question_id': 'cx-4',
            },
            {
                'id': 'cx-4',
                'prompt': 'Are you looking for information on storage and sharing of Contractual documents ?',
                'guidance': 'Management of Contractual Documents must be via approved methods as described in this section.',
                'answer': None,
                'yes_next_question_id': 'cx-8',
                'no_next_question_id': 'cx-5',
            },
            {
                'id': 'cx-5',
                'prompt': 'Are you looking for information on storage and sharing of Supporting Material ?',
                'guidance': 'Management of Supporting Materials must be via approved methods as described in this section.',
                'answer': None,
                'yes_next_question_id': 'cx-10',
                'no_next_question_id': 'cx-12',
            },
            {
                'id': 'cx-6',
                'prompt': 'Are you looking for information about storage of Deliverables ?',
                'guidance': 'This section provides guidance on approved methods for storage of Deliverables.',
                'answer': None,
                'yes_outcome_id': 'cx-o13',
                'no_next_question_id': 'cx-7',
            },
            {
                'id': 'cx-7',
                'prompt': 'Are you looking for information about sharing of Deliverables ?',
                'guidance': 'This section provides guidance on approved methods for sharing of Deliverables.',
                'answer': None,
                'yes_outcome_id': 'cx-o14',
                'no_outcome_id': 'cx-o20',
            },
            {
                'id': 'cx-8',
                'prompt': 'Are you looking for information on storage of Contractual documents ?',
                'guidance': 'This section provides guidance on approved methods for storage of Contractual Documents.',
                'answer': None,
                'yes_outcome_id': 'cx-o15',
                'no_next_question_id': 'cx-9',
            },
            {
                'id': 'cx-9',
                'prompt': 'Are you looking for information on sharing of Contractual documents ?',
                'guidance': 'This section provides guidance on approved methods for sharing of Contractual Documents.',
                'answer': None,
                'yes_outcome_id': 'cx-o16',
                'no_outcome_id': 'cx-o20',
            },
            {
                'id': 'cx-10',
                'prompt': 'Are you looking for information on storage of Supporting Material ?',
                'guidance': 'This section provides guidance on approved methods for storage of Supporting Material.',
                'answer': None,
                'yes_outcome_id': 'cx-o17',
                'no_next_question_id': 'cx-11',
            },
            {
                'id': 'cx-11',
                'prompt': 'Are you looking for information on sharing of Supporting Material ?',
                'guidance': 'This section provides guidance on approved methods for sharing of Supporting Material.',
                'answer': None,
                'yes_outcome_id': 'cx-o18',
                'no_outcome_id': 'cx-o20',
            },
            {
                'id': 'cx-12',
                'prompt': 'Are you looking for information on the exception process ?',
                'guidance': 'This section provides guidance on the process to request exceptions to the Customer Experience Professional Services Document Management Policy.',
                'answer': None,
                'yes_outcome_id': 'cx-o19',
                'no_outcome_id': 'cx-o20',
            },
        ],
        'outcomes': [
            {
                'id': 'cx-o13',
                'title': 'Storage of Deliverables',
                'details_html': (
                    '<p>All Deliverables are <strong>sensitive information</strong> by default without exception.</p>'
                    '<table><tbody>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/601/">DCP</a></td><td>Final peer reviewed Deliverables must be stored in DCP before internal approval and customer acceptance. Drafts may also be stored there. <strong>Draft: Yes. Final: Yes - Mandatory.</strong></td></tr>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/600/">Doc Central</a></td><td>Can be used for storage and sharing, including Doc Exchange integration. <strong>Yes - Copy.</strong></td></tr>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/529/">Doc Exchange</a></td><td>Use for short-term sharing; move retained files to Doc Central or DCP. <strong>Yes - Copy.</strong></td></tr>'
                    '<tr><td><a href="https://scdp.cisco.com/">SCDP</a></td><td>Drafts and copies can live in SCDP, but final versions must still be uploaded to DCP. <strong>Draft: Yes. Final: Yes - Copy, DCP mandatory.</strong></td></tr>'
                    '<tr><td><a href="https://teams.webex.com/signin">Cisco Webex Teams</a></td><td>Allowed for collaboration while in draft and review stages. Final reviewed versions must move to DCP. <strong>Draft: Yes. Final: No.</strong></td></tr>'
                    '<tr><td>Office 365 (SharePoint, OneDrive)</td><td>Drafts and copies can be stored, but final reviewed versions must move to DCP. <strong>Draft: Yes. Final: Yes - Copy, DCP mandatory.</strong></td></tr>'
                    '<tr><td>Office 365 (OneNote)</td><td>Not approved for deliverables. <strong>Draft: No. Final: No.</strong></td></tr>'
                    '<tr><td>All Other Tools</td><td>No other tools are approved for use.</td></tr>'
                    '</tbody></table>'
                    '<p><strong>What is Sensitive Information?</strong> Information that, if exposed, could compromise Cisco or Customer operations or reputation.</p>'
                ),
            },
            {
                'id': 'cx-o14',
                'title': 'Sharing of Deliverables',
                'details_html': (
                    '<p>All Deliverables are <strong>sensitive information</strong> by default without exception.</p>'
                    '<table><tbody>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/601/">DCP</a></td><td>Preferred internal and external sharing method for draft and final reviewed deliverables. <strong>Internal: Yes. External: Yes - Preferred.</strong></td></tr>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/600/">Doc Central</a></td><td>Can share internally and externally through Doc Exchange integration. <strong>Internal: Yes - Copy. External: Yes - Copy.</strong></td></tr>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/529/">Doc Exchange</a></td><td>Use for short-term internal and external sharing. <strong>Internal: Yes - Copy. External: Yes - Copy.</strong></td></tr>'
                    '<tr><td><a href="https://scdp.cisco.com/">SCDP</a></td><td>Good for internal collaboration only; final versions must still go to DCP. <strong>Internal: Yes. External: No.</strong></td></tr>'
                    '<tr><td><a href="https://dds.cisco.com">DDS</a></td><td>Designed for draft sharing and customer feedback for documents authored in SCDP. <strong>Internal: Yes - Draft only. External: Yes - Draft only.</strong></td></tr>'
                    '<tr><td><a href="https://teams.webex.com/signin">Cisco Webex Teams</a></td><td>Can be used while in draft, peer review and customer review stages. <strong>Internal: Yes. External: Yes.</strong></td></tr>'
                    '<tr><td>Office 365 (SharePoint)</td><td>Can be used for collaboration and copies of final reviewed versions. <strong>Internal: Yes. External: Yes.</strong></td></tr>'
                    '<tr><td>Office 365 (OneDrive)</td><td>Internal collaboration only. <strong>Internal: Yes. External: No.</strong></td></tr>'
                    '<tr><td>Email / File Transfer Services / SW Center</td><td>Encrypted email and approved transfer tools are allowed as described in policy. Unencrypted email is not approved for external sensitive sharing.</td></tr>'
                    '<tr><td>All Other Tools</td><td>No other tools are approved for use.</td></tr>'
                    '</tbody></table>'
                ),
            },
            {
                'id': 'cx-o15',
                'title': 'Storage of Contractual Documents',
                'details_html': (
                    '<p>All Contractual Documents are <strong>sensitive information</strong> by default without exception.</p>'
                    '<table><tbody>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/601/">DCP</a></td><td>DCP is the master location for Contractual Information. <strong>Yes.</strong></td></tr>'
                    '<tr><td><a href="https://scdp.cisco.com/">SCDP</a></td><td>Drafts and copies may be stored for internal collaboration, but final versions must be in DCP. <strong>Draft: Yes. Final: Yes - Copy, DCP mandatory.</strong></td></tr>'
                    '<tr><td><a href="https://teams.webex.com/signin">Cisco Webex Teams</a></td><td>Draft collaboration only; final versions must be in DCP. <strong>Draft: Yes. Final: No.</strong></td></tr>'
                    '<tr><td>Office 365 (SharePoint, OneDrive)</td><td>Drafts and copies can be stored for collaboration. Final versions must be in DCP. <strong>Draft: Yes. Final: No.</strong></td></tr>'
                    '<tr><td>All Other Tools</td><td>No other tools are approved for use.</td></tr>'
                    '</tbody></table>'
                ),
            },
            {
                'id': 'cx-o16',
                'title': 'Sharing of Contractual Documents',
                'details_html': (
                    '<p>All Contractual Documents are <strong>sensitive information</strong> by default without exception.</p>'
                    '<table><tbody>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/601/">DCP</a></td><td>Master location for internal and external sharing. <strong>Internal: Yes. External: Yes.</strong></td></tr>'
                    '<tr><td><a href="https://scdp.cisco.com/">SCDP</a></td><td>Internal collaboration only. <strong>Internal: Yes. External: No.</strong></td></tr>'
                    '<tr><td><a href="https://teams.webex.com/signin">Cisco Webex Teams</a></td><td>Internal collaboration only for contractual docs. <strong>Internal: Yes. External: No.</strong></td></tr>'
                    '<tr><td>Office 365 (SharePoint)</td><td>Copies of final versions can be shared internally and externally if access is managed. <strong>Internal: Yes - Copy. External: Yes - Copy.</strong></td></tr>'
                    '<tr><td>Office 365 (OneDrive)</td><td>Internal collaboration only. <strong>Internal: Yes - Copy. External: No.</strong></td></tr>'
                    '<tr><td>Email (encrypted)</td><td>Allowed for internal and external sharing.</td></tr>'
                    '<tr><td>Email (unencrypted)</td><td>Internal only for sensitive information. External is not approved.</td></tr>'
                    '<tr><td>All Other Tools</td><td>No other tools are approved for use.</td></tr>'
                    '</tbody></table>'
                ),
            },
            {
                'id': 'cx-o17',
                'title': 'Storage of Supporting Material',
                'details_html': (
                    '<table><tbody>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/601/">DCP</a></td><td>Preferred location alongside Customer Deliverables. <strong>Non-sensitive: Yes. Sensitive: Yes.</strong></td></tr>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/600/">Doc Central</a></td><td>Can be used for storage and sharing. <strong>Non-sensitive: Yes. Sensitive: Yes.</strong></td></tr>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/529/">Doc Exchange</a></td><td>Short-term storage for sharing. <strong>Non-sensitive: Yes. Sensitive: Yes.</strong></td></tr>'
                    '<tr><td><a href="https://scdp.cisco.com/">SCDP</a></td><td>Internal storage for non-sensitive and sensitive drafts. Final retained material should move to an approved repository. <strong>Draft: Yes. Final: No.</strong></td></tr>'
                    '<tr><td><a href="https://teams.webex.com/signin">Cisco Webex Teams</a></td><td>Collaboration on draft supporting material. Retained records must move to DCP or Doc Central. <strong>Draft: Yes. Final: No.</strong></td></tr>'
                    '<tr><td>Office 365 / JIRA</td><td>Can be used depending on sensitivity and record-retention requirements as summarized in policy.</td></tr>'
                    '<tr><td>All Other Tools</td><td>No other tools are approved for use.</td></tr>'
                    '</tbody></table>'
                    '<p><strong>What is Sensitive Information?</strong> Most supporting material is non-sensitive, but PMs are responsible for excluding, cross-referencing or redacting sensitive content before using lower-control tools.</p>'
                ),
            },
            {
                'id': 'cx-o18',
                'title': 'Sharing of Supporting Material',
                'details_html': (
                    '<table><tbody>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/601/">DCP</a></td><td>Preferred location. <strong>Internal & External, Non-sensitive: Yes. Sensitive: Yes.</strong></td></tr>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/600/">Doc Central</a></td><td>Allowed for internal and external sharing. <strong>Non-sensitive: Yes. Sensitive: Yes.</strong></td></tr>'
                    '<tr><td><a href="https://dataadvisor.cloudapps.cisco.com/da/tools/529/">Doc Exchange</a></td><td>Allowed for internal and external sharing. <strong>Non-sensitive: Yes. Sensitive: Yes.</strong></td></tr>'
                    '<tr><td><a href="https://scdp.cisco.com/">SCDP</a></td><td>Internal sharing only, not for external or sensitive external use. <strong>Internal Non-sensitive: Yes. Sensitive: No. External: No.</strong></td></tr>'
                    '<tr><td><a href="https://teams.webex.com/signin">Cisco Webex Teams</a></td><td>Allowed for internal and external collaboration up to Highly Confidential. <strong>Internal & External, Non-sensitive: Yes. Sensitive: Yes.</strong></td></tr>'
                    '<tr><td>Office 365 / Email / JIRA</td><td>Use according to sensitivity and retention rules described in policy.</td></tr>'
                    '<tr><td>All Other Tools</td><td>No other tools are approved for use.</td></tr>'
                    '</tbody></table>'
                ),
            },
            {
                'id': 'cx-o19',
                'title': 'Exception Process',
                'details_html': (
                    '<p>The exception process caters for customer-initiated requests for alternate tool use.</p>'
                    '<ol>'
                    '<li>Customer requests an exception and the PM explains data protection risks.</li>'
                    '<li>PM emails the Risk Acceptance Disclaimer text to the customer.</li>'
                    '<li>Customer returns the completed disclaimer and the PM uploads it to DCP.</li>'
                    '<li>PM submits the exception request via the GSM Policy & Process Governance mailer.</li>'
                    '<li>Forum review is initiated and supporting evidence is uploaded for review.</li>'
                    '<li>Decision is communicated, recorded in DCP, and then shared internally and externally.</li>'
                    '<li>Subscription projects require annual review with the customer.</li>'
                    '</ol>'
                    '<p>Upload evidence to DCP as <strong>Supporting Material</strong> with keyword <strong>doc_sharing_exception</strong>.</p>'
                ),
            },
            {
                'id': 'cx-o20',
                'title': 'Wrong Advisor',
                'details_html': '<p>It seems you may have the wrong advisor. This advisor only covers guidance for tools, as they pertain to CX Professional Services deliverables, contractual information, or supporting materials.</p>',
            },
        ],
    },
    {
        'id': 'bookings',
        'name': 'Bookings',
        'description': (
            'Data classification decision tree for Bookings data or associated reporting. If an '
            'application, database, report, or other deliverable contains data with various '
            'classification levels, the highest classification level is applied as the minimum '
            'classification for the grouped data.'
        ),
        'max_questions': 8,
        'questions': [
            {
                'id': 'bookings-1',
                'prompt': 'Includes current quarter Bookings actuals or future (ie., forecast) data?',
                'guidance': (
                    'Current quarter actuals and future forecast data are sensitive until they are '
                    'publicly shared.'
                ),
                'answer': 'YES',
            },
            {
                'id': 'bookings-2',
                'prompt': 'Includes Product data?',
                'guidance': '',
                'answer': 'NO',
            },
            {
                'id': 'bookings-3',
                'prompt': 'Includes Cost data?',
                'guidance': '',
                'answer': None,
            },
        ],
    },
    {
        'id': 'conflict-board-membership',
        'name': 'Conflict of Interest - Board Membership',
        'description': (
            'The Board Membership decision tree provides guidance to establish if a disclosure is '
            'required if you plan on serving on a government affiliated board, Industry '
            'Association board, Non-profit organization board or outside board of directors of a '
            'for-profit organization. Please select the responses specific to your scenario to '
            'receive the relevant guidance.'
        ),
        'max_questions': 9,
        'questions': [],
    },
    {
        'id': 'conflict-family-member-employment',
        'name': 'Conflict of Interest - Family Member Employment',
        'description': (
            'Conflict of Interest - The Family Member Employment decision tree provides guidance '
            'to determine if a disclosure is required for a family member that works at Cisco (or '
            'is in the process of being hired) or holds a position with a Cisco partner, customer, '
            'supplier, competitor or acquisition candidate. Please select the responses specific '
            'to your scenario to receive the relevant guidance.'
        ),
        'max_questions': 5,
        'questions': [],
    },
    {'id': 'customer-success-assets', 'name': 'Customer Success Assets', 'description': 'Guidance on tool usage for customer success assets and success plan documentation.', 'max_questions': 7, 'questions': []},
    {'id': 'deal-approval-data', 'name': 'Deal Approval Data', 'description': 'Decision support for deal approval files, attachments and approval notes.', 'max_questions': 6, 'questions': []},
    {'id': 'delivery-governance', 'name': 'Delivery Governance', 'description': 'Guidance for governance notes, project playbooks and delivery stage documents.', 'max_questions': 10, 'questions': []},
    {'id': 'finance-operations', 'name': 'Finance Operations', 'description': 'Advisor for finance operations content and internal reporting packages.', 'max_questions': 6, 'questions': []},
    {'id': 'hr-sensitive-content', 'name': 'HR Sensitive Content', 'description': 'Guidance for storing and sharing human resources related documents.', 'max_questions': 7, 'questions': []},
    {'id': 'install-base-exports', 'name': 'Install Base Exports', 'description': 'Classification help for install base extracts and asset related exports.', 'max_questions': 8, 'questions': []},
    {'id': 'partner-programs', 'name': 'Partner Programs', 'description': 'Decision tree for partner engagement deliverables and incentive reports.', 'max_questions': 7, 'questions': []},
    {'id': 'procurement', 'name': 'Procurement', 'description': 'Guidance for procurement documents, sourcing packs and supplier materials.', 'max_questions': 8, 'questions': []},
    {'id': 'product-launch', 'name': 'Product Launch', 'description': 'Classification advisor for launch plans, launch scorecards and launch assets.', 'max_questions': 5, 'questions': []},
    {'id': 'sales-compensation', 'name': 'Sales Compensation', 'description': 'Guidance for quota, incentive and compensation planning content.', 'max_questions': 6, 'questions': []},
    {'id': 'security-review', 'name': 'Security Review', 'description': 'Decision support for security review packages and evidence repositories.', 'max_questions': 7, 'questions': []},
    {'id': 'strategy-planning', 'name': 'Strategy Planning', 'description': 'Guidance for strategy decks, long-range planning and internal positioning.', 'max_questions': 8, 'questions': []},
    {'id': 'support-escalations', 'name': 'Support Escalations', 'description': 'Classification advisor for support cases, escalations and root-cause notes.', 'max_questions': 5, 'questions': []},
]


_apply_cx_workbook_overrides(ADVISOR_CATALOG[0])


def list_advisors(session: Session) -> list[AdvisorRead]:
    answers = _answer_lookup(session)
    return [_hydrate_advisor(advisor, answers) for advisor in ADVISOR_CATALOG]


def get_advisor(session: Session, advisor_id: str) -> AdvisorRead:
    advisor = next((item for item in ADVISOR_CATALOG if item['id'] == advisor_id), None)
    if advisor is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Advisor not found.')

    return _hydrate_advisor(advisor, _answer_lookup(session))


def update_answer(session: Session, advisor_id: str, question_id: str, answer: str) -> AdvisorRead:
    advisor = next((item for item in ADVISOR_CATALOG if item['id'] == advisor_id), None)
    if advisor is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Advisor not found.')

    questions = advisor.get('questions', [])
    if not any(question['id'] == question_id for question in questions):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Question not found.')

    existing = session.scalar(
        select(AdvisorAnswer).where(
            AdvisorAnswer.advisor_id == advisor_id,
            AdvisorAnswer.question_id == question_id,
        )
    )

    if existing is None:
        existing = AdvisorAnswer(advisor_id=advisor_id, question_id=question_id, answer=answer)
    else:
        existing.answer = answer

    session.add(existing)
    session.commit()
    return get_advisor(session, advisor_id)


def reset_answers(session: Session, advisor_id: str) -> AdvisorRead:
    advisor = next((item for item in ADVISOR_CATALOG if item['id'] == advisor_id), None)
    if advisor is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Advisor not found.')

    session.execute(delete(AdvisorAnswer).where(AdvisorAnswer.advisor_id == advisor_id))
    session.commit()
    return get_advisor(session, advisor_id)


def _answer_lookup(session: Session) -> dict[tuple[str, str], str]:
    rows = session.scalars(select(AdvisorAnswer)).all()
    return {(row.advisor_id, row.question_id): row.answer for row in rows}


def _hydrate_advisor(advisor: dict[str, object], answers: dict[tuple[str, str], str]) -> AdvisorRead:
    payload = deepcopy(advisor)
    questions = payload.get('questions', [])
    for question in questions:
        key = (str(payload['id']), str(question['id']))
        if key in answers:
            question['answer'] = answers[key]

    current_question_id, current_outcome_id, answered_count = _resolve_progress(payload)
    payload['current_question_id'] = current_question_id
    payload['current_outcome_id'] = current_outcome_id
    payload['answered_count'] = answered_count
    return AdvisorRead.model_validate(payload)


def _resolve_progress(advisor: dict[str, object]) -> tuple[str | None, str | None, int]:
    questions = advisor.get('questions', [])
    if not questions:
        return None, None, 0

    question_lookup = {question['id']: question for question in questions}
    current_id = questions[0]['id']
    answered_count = 0

    while current_id is not None:
        question = question_lookup[current_id]
        answer = question.get('answer')
        if answer not in {'YES', 'NO'}:
            return current_id, None, answered_count

        answered_count += 1
        if answer == 'YES':
            next_question = question.get('yes_next_question_id')
            next_outcome = question.get('yes_outcome_id')
        else:
            next_question = question.get('no_next_question_id')
            next_outcome = question.get('no_outcome_id')

        if next_outcome:
            return None, str(next_outcome), answered_count

        current_id = str(next_question) if next_question else None

    first_unanswered = next(
        (question['id'] for question in questions if question.get('answer') not in {'YES', 'NO'}),
        None,
    )
    return first_unanswered, None, answered_count