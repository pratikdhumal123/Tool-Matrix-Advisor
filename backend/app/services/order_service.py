from __future__ import annotations

import csv
import io
import random
from datetime import datetime, timezone

from fastapi import HTTPException, status
from openpyxl import load_workbook
from sqlalchemy import desc, select
from sqlalchemy.orm import Session

from app.models.order import Order, OrderStatus
from app.schemas.order import DashboardResponse, DashboardSummary, OrderCreate


SAMPLE_ORDERS = [
    {
        'order_ref': 'VOKGLBMK317191',
        'customer_name': 'Aanya Singh',
        'advisor_name': 'Rohan',
        'counter_name': '(Club House) Global Counter - New',
        'counter_code': '3173',
        'item_name': 'Protein bowl',
        'quantity': 1,
        'total_amount': 460,
        'status': OrderStatus.READY,
        'pin': '471430',
    },
    {
        'order_ref': 'MORADV-2026-1002',
        'customer_name': 'Kabir Mehta',
        'advisor_name': 'Rohan',
        'counter_name': 'Fitness Juice Bar',
        'counter_code': '4201',
        'item_name': 'Hydration shake',
        'quantity': 2,
        'total_amount': 780,
        'status': OrderStatus.READY,
        'pin': '512884',
    },
    {
        'order_ref': 'MORADV-2026-1003',
        'customer_name': 'Sara Thomas',
        'advisor_name': 'Neha',
        'counter_name': 'Wellness Kitchen',
        'counter_code': '2214',
        'item_name': 'Salad combo',
        'quantity': 1,
        'total_amount': 540,
        'status': OrderStatus.PENDING,
        'pin': None,
    },
    {
        'order_ref': 'MORADV-2026-1004',
        'customer_name': 'Vihaan Patel',
        'advisor_name': 'Neha',
        'counter_name': 'Runner\'s Counter',
        'counter_code': '1820',
        'item_name': 'Energy bar pack',
        'quantity': 3,
        'total_amount': 910,
        'status': OrderStatus.COLLECTED,
        'pin': '188204',
    },
]

HEADER_ALIASES = {
    'order ref': 'order_ref',
    'order_ref': 'order_ref',
    'customer': 'customer_name',
    'customer name': 'customer_name',
    'customer_name': 'customer_name',
    'advisor': 'advisor_name',
    'advisor name': 'advisor_name',
    'advisor_name': 'advisor_name',
    'counter': 'counter_name',
    'counter name': 'counter_name',
    'counter_name': 'counter_name',
    'counter code': 'counter_code',
    'counter_code': 'counter_code',
    'item': 'item_name',
    'item name': 'item_name',
    'item_name': 'item_name',
    'quantity': 'quantity',
    'total': 'total_amount',
    'amount': 'total_amount',
    'total_amount': 'total_amount',
    'status': 'status',
}


def seed_orders(session: Session) -> None:
    existing = session.scalar(select(Order.id).limit(1))
    if existing is not None:
        return

    now = datetime.now(timezone.utc)
    for index, payload in enumerate(SAMPLE_ORDERS):
        order = Order(
            **payload,
            qr_payload=(
                f"advisor-sale:{payload['order_ref']}:{payload['pin']}" if payload['pin'] else None
            ),
            qr_generated_at=now if payload['pin'] else None,
            created_at=now,
            collected_at=now if payload['status'] == OrderStatus.COLLECTED else None,
        )
        order.created_at = now.replace(minute=max(now.minute - index, 0))
        session.add(order)

    session.commit()


def list_orders(session: Session) -> list[Order]:
    statement = select(Order).order_by(desc(Order.created_at))
    return list(session.scalars(statement))


def create_order(session: Session, payload: OrderCreate) -> Order:
    order = Order(**payload.model_dump())
    session.add(order)
    session.commit()
    session.refresh(order)
    return order


def build_dashboard(session: Session) -> DashboardResponse:
    orders = list_orders(session)
    ready_orders = [order for order in orders if order.status == OrderStatus.READY]
    pending_orders = [order for order in orders if order.status == OrderStatus.PENDING]
    collected_orders = [order for order in orders if order.status == OrderStatus.COLLECTED]
    total = len(orders)
    collection_rate = int((len(collected_orders) / total) * 100) if total else 0
    revenue_in_queue = sum(order.total_amount for order in ready_orders + pending_orders)

    summary = DashboardSummary(
        total_orders=total,
        ready_orders=len(ready_orders),
        pending_orders=len(pending_orders),
        collected_orders=len(collected_orders),
        collection_rate=collection_rate,
        revenue_in_queue=revenue_in_queue,
    )
    return DashboardResponse(
        summary=summary,
        spotlight=ready_orders[:3],
        recent_collections=collected_orders[:3],
    )


def generate_collection(session: Session, order_id: int) -> Order:
    order = session.get(Order, order_id)
    if order is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Order not found.')

    if not order.pin:
        order.pin = f'{random.randint(0, 999999):06d}'

    order.status = OrderStatus.READY
    order.qr_payload = f'advisor-sale:{order.order_ref}:{order.pin}'
    order.qr_generated_at = datetime.now(timezone.utc)
    session.add(order)
    session.commit()
    session.refresh(order)
    return order


def collect_by_pin(session: Session, pin: str) -> Order:
    statement = select(Order).where(Order.pin == pin)
    order = session.scalar(statement)
    if order is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail='Collection pin not found.')
    if order.status == OrderStatus.COLLECTED:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail='Order is already collected.')

    order.status = OrderStatus.COLLECTED
    order.collected_at = datetime.now(timezone.utc)
    session.add(order)
    session.commit()
    session.refresh(order)
    return order


def import_orders(session: Session, filename: str, content: bytes) -> tuple[int, int, list[str]]:
    rows = _read_rows(filename, content)
    imported = 0
    updated = 0
    errors: list[str] = []

    for row_number, row in enumerate(rows, start=2):
        try:
            payload = _normalize_row(row)
        except ValueError as exc:
            errors.append(f'Row {row_number}: {exc}')
            continue

        existing = session.scalar(select(Order).where(Order.order_ref == payload.order_ref))
        if existing is None:
            session.add(Order(**payload.model_dump()))
            imported += 1
            continue

        for key, value in payload.model_dump().items():
            setattr(existing, key, value)
        updated += 1

    session.commit()
    return imported, updated, errors


def _read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    lower_name = filename.lower()
    if lower_name.endswith('.csv'):
        text = content.decode('utf-8-sig')
        return list(csv.DictReader(io.StringIO(text)))

    if lower_name.endswith('.xlsx'):
        workbook = load_workbook(io.BytesIO(content), data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(value).strip() if value is not None else '' for value in rows[0]]
        parsed_rows: list[dict[str, str]] = []
        for values in rows[1:]:
            parsed_rows.append(
                {
                    headers[index]: '' if value is None else str(value).strip()
                    for index, value in enumerate(values)
                    if index < len(headers)
                }
            )
        return parsed_rows

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail='Unsupported file type. Upload .csv or .xlsx files.',
    )


def _normalize_row(row: dict[str, str]) -> OrderCreate:
    normalized: dict[str, str] = {}
    for key, value in row.items():
        mapped_key = HEADER_ALIASES.get(key.strip().lower())
        if mapped_key:
            normalized[mapped_key] = value.strip() if isinstance(value, str) else value

    required = [
        'order_ref',
        'customer_name',
        'advisor_name',
        'counter_name',
        'counter_code',
        'item_name',
    ]
    missing = [field for field in required if not normalized.get(field)]
    if missing:
        raise ValueError(f"Missing required values for: {', '.join(missing)}")

    normalized['quantity'] = int(normalized.get('quantity') or 1)
    normalized['total_amount'] = float(normalized.get('total_amount') or 0)
    normalized['status'] = str(normalized.get('status') or OrderStatus.PENDING).upper()

    return OrderCreate(**normalized)